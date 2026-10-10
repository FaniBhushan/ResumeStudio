from __future__ import annotations

from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader


MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_EXTRACTED_CHARACTERS = 120_000


def extract_cv_text(filename: str, data: bytes) -> str:
    """Extract selectable text from a PDF, DOCX, or plain-text upload."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("The file is larger than 20 MB. Please upload a smaller CV.")
    extension = Path(filename).suffix.casefold()
    if extension == ".pdf":
        try:
            reader = PdfReader(BytesIO(data), strict=False)
            text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:
            raise ValueError("Could not read this PDF. Please check that it opens correctly and try again.") from exc
        if not text.strip():
            raise ValueError("This PDF has no selectable text. Scanned PDFs need OCR, which is not supported yet.")
    elif extension == ".docx":
        try:
            document = Document(BytesIO(data))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
                    if any(cells):
                        parts.append(" | ".join(cells))
            text = "\n".join(parts)
        except Exception as exc:
            raise ValueError("Could not read this Word document. Please upload a valid .docx file.") from exc
    elif extension == ".txt":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = data.decode("cp1252")
            except UnicodeDecodeError as exc:
                raise ValueError("Could not decode this text file. Please save it as UTF-8 and try again.") from exc
    else:
        raise ValueError("Use a text-based PDF, a .docx Word document, or a .txt file.")

    text = text.strip()
    if not text:
        raise ValueError("No readable text was found in this file.")
    if len(text) > MAX_EXTRACTED_CHARACTERS:
        raise ValueError("The extracted CV is unusually long. Please shorten the document and try again.")
    return text


def extract_job_description(filename: str, data: bytes) -> str:
    """Extract and validate a job description from a supported upload."""
    if Path(filename).suffix.casefold() == ".md":
        if len(data) > MAX_UPLOAD_BYTES:
            raise ValueError("The file is larger than 20 MB. Please upload a smaller job description.")
        try:
            text = data.decode("utf-8-sig").strip()
        except UnicodeDecodeError as exc:
            raise ValueError("Could not decode this Markdown file. Please save it as UTF-8 and try again.") from exc
        if not text:
            raise ValueError("No readable text was found in this job description file.")
        if len(text) > 50_000:
            raise ValueError("The job description is too long. Please limit it to 50,000 characters.")
        return text

    text = extract_cv_text(filename, data)
    if len(text) > 50_000:
        raise ValueError("The job description is too long. Please limit it to 50,000 characters.")
    return text
